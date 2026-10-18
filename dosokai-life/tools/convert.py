#!/usr/bin/env python3
"""events.xlsx を、ゲームが読む data/events.json と data/events.js に変換する。

使い方:
    python3 tools/convert.py                       # data/events.xlsx → data/events.js
    python3 tools/convert.py 入力.xlsx 出力.js       # 場所を指定する場合

events.js は events.json の中身を window.GAME_DATA に入れただけのもの。
ブラウザはローカルの JSON を fetch で読めないため、<script> で読める形も作る。
条件・効果の書式の確認はゲーム側（js/rules.js）が行い、読めない行は画面の隅に出る。
"""
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl が必要です: pip install openpyxl')

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)


def text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace('\r\n', '\n').strip()


def num(v, default=None):
    if v is None or text(v) == '':
        return default
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return text(v)  # 数字でなければ文字のまま渡し、ゲーム側で警告する


def sheet_rows(wb, name):
    """1行目を見出しとして、見出し名 → 値 の辞書を行ごとに返す（列の順番に依存しない）。シートがなければ空。"""
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    head = [text(h) for h in rows[0]]
    out = []
    for i, r in enumerate(rows[1:], start=2):
        d = {h: r[j] if j < len(r) else None for j, h in enumerate(head) if h}
        d['_row'] = i
        out.append(d)
    return out


def convert_events(wb, sheet='イベント'):
    events = []
    for d in sheet_rows(wb, sheet):
        eid = text(d.get('ID'))
        if not eid or eid == 'メモ':
            continue
        choices = []
        for n in range(1, 9):  # 選択肢は8つまで
            label = text(d.get(f'選択肢{n}'))
            if label:
                choices.append({
                    'label': label,
                    'effect': text(d.get(f'効果{n}')),
                    'cond': text(d.get(f'選択肢{n}の条件')),
                })
        events.append({
            'id': eid,
            'row': d['_row'],
            'type': text(d.get('種類')),
            'ageFrom': num(d.get('年齢(から)'), 18),
            'ageTo': num(d.get('年齢(まで)'), 35),
            'cond': text(d.get('条件')),
            'prob': num(d.get('確率'), 0),
            'once': text(d.get('一度だけ')) in ('○', '〇', 'o', 'O', '1'),
            'title': text(d.get('タイトル')),
            'text': text(d.get('本文')),
            'choices': choices,
            'effect': text(d.get('効果(選択肢なし)')),
            'probMod': text(d.get('確率の補正')),
            'staging': text(d.get('演出')),
            'category': text(d.get('分類')),
            # 現実でそれぞれの選択肢を選ぶ人の割合（照合に使う。空なら「ふつうの選び方」）
            'weights': [float(x) for x in text(d.get('選ばれやすさ')).replace('、', ',').split(',') if x.strip()],
        })
        if sheet == 'くじ':
            events[-1].update({
                'group': text(d.get('グループ')),
                'rarity': num(d.get('レア度'), 0),
                'lifetime': text(d.get('一生で')),
                'confidence': text(d.get('確度')),
                'basis': text(d.get('根拠')),
            })
    return events


def convert_reunion(wb):
    out = []
    for d in sheet_rows(wb, '同窓会'):
        rid = text(d.get('ID'))
        if not rid or rid == 'メモ':
            continue
        out.append({
            'id': rid,
            'age': num(d.get('年齢')),
            'scene': text(d.get('場面')),
            'cond': text(d.get('条件')),
            'text': text(d.get('本文')),
            'staging': text(d.get('演出')),
        })
    return out


def convert_classmates(wb):
    out = {}
    for d in sheet_rows(wb, '同級生'):
        name = text(d.get('名前'))
        if not name:
            continue
        c = out.setdefault(name, {'name': name, 'profile': '', 'patterns': []})
        if text(d.get('人物')):
            c['profile'] = text(d.get('人物'))
        c['patterns'].append({
            'key': text(d.get('パターン')),
            'title': text(d.get('運命')),
            'prob': num(d.get('確率'), 0),
            'at25': text(d.get('25歳の様子')),
            'at35': text(d.get('35歳の様子')),
            'job25': text(d.get('25歳の仕事')),
            'job35': text(d.get('35歳の仕事')),
            'at50': text(d.get('50歳の様子')),
            'at65': text(d.get('65歳の様子')),
            'job50': text(d.get('50歳の仕事')),
            'job65': text(d.get('65歳の仕事')),
            'mortality': num(d.get('亡くなりやすさ'), 1),
        })
    return list(out.values())


def convert_careers(wb):
    out = []
    for d in sheet_rows(wb, 'キャリア'):
        name = text(d.get('キャリア'))
        if not name:
            continue
        out.append({
            'name': name,
            'how': text(d.get('なり方')),
            'yearly': text(d.get('毎年の変化')),
            'strong': text(d.get('得意な数値')),
            'weak': text(d.get('削れやすい数値')),
            'fortunes': text(d.get('会社の運命（仮の確率）')),
            'kind': text(d.get('区分')),
            'industry': text(d.get('業種')),
            'shareUni': num(d.get('大卒の割合')),
            'shareHigh': num(d.get('高卒の割合')),
            'uniOnly': text(d.get('大卒のみ')) in ('○', '〇', '1'),
            'pass': num(d.get('合格率')),
            'econ': text(d.get('景気の影響')),
            'heading': text(d.get('求人の見出し')),
            'faculties': text(d.get('学部')),
            'female': num(d.get('女性の割合')),
            'examples': [x.strip() for x in text(d.get('会社の例')).replace('、', ',').split(',') if x.strip()],
            'source': text(d.get('出典')),
        })
    return out


def convert_simple(wb, sheet, columns):
    """見出し → キー の対応で、そのまま行を写す（数値は数値に）。"""
    out = []
    for d in sheet_rows(wb, sheet):
        if not any(text(d.get(h)) for h in columns):
            continue
        row = {'row': d['_row']}
        for head, (key, kind) in columns.items():
            row[key] = num(d.get(head)) if kind == 'num' else text(d.get(head))
        out.append(row)
    return out


def list_assets():
    """assets フォルダにある絵の一覧（画面は、ここにある絵だけを使う。ない絵は読みに行かない）。"""
    base = os.path.join(ROOT, 'assets')
    out = []
    for d, _, files in os.walk(base):
        for f in files:
            if f.lower().endswith(('.webp', '.png', '.jpg', '.jpeg', '.svg')):
                out.append(os.path.relpath(os.path.join(d, f), ROOT).replace(os.sep, '/'))
    return sorted(out)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'data', 'events.xlsx')
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'data', 'events.js')
    wb = openpyxl.load_workbook(src, data_only=True)
    data = {
        'source': os.path.basename(src),
        'events': convert_events(wb),
        'reunion': convert_reunion(wb),
        'classmates': convert_classmates(wb),
        'careers': convert_careers(wb),
        'tables': convert_simple(wb, '確率表', {
            '名前': ('name', 'text'), '基準': ('basis', 'text'), 'から': ('from', 'num'),
            'まで': ('to', 'num'), '年率': ('rate', 'num'), '出典': ('source', 'text')}),
        'stats': convert_simple(wb, '統計', {
            '指標': ('name', 'text'), '測り方': ('measure', 'text'), '現実の値': ('real', 'num'),
            '下限': ('low', 'num'), '上限': ('high', 'num'), '出典': ('source', 'text')}),
        'partner': convert_simple(wb, '配偶者', {
            '属性': ('attr', 'text'), '候補': ('value', 'text'), '割合': ('weight', 'num'),
            '条件': ('cond', 'text'), '値': ('num', 'num'), '文章': ('text', 'text'), '出典': ('source', 'text')}),
        'people': convert_simple(wb, '人物', {
            '人物': ('id', 'text'), '表示名': ('label', 'text'), 'ファイル名の形': ('file', 'text'),
            '年代': ('ages', 'text'), '表情': ('expressions', 'text'), '絵の説明': ('description', 'text')}),
        'rare': convert_events(wb, 'くじ'),
        'assets': list_assets(),
    }
    body = json.dumps(data, ensure_ascii=False, indent=1)
    with open(os.path.splitext(dst)[0] + '.json', 'w', encoding='utf-8') as f:
        f.write(body + '\n')
    with open(dst, 'w', encoding='utf-8') as f:
        f.write('// tools/convert.py が events.xlsx から作ったファイルです。手で編集しないでください。\n')
        f.write('window.GAME_DATA = ' + body + ';\n')
    print(f"変換しました: イベント {len(data['events'])} 件 / くじ {len(data['rare'])} 件 / "
          f"同窓会 {len(data['reunion'])} 件 / 同級生 {len(data['classmates'])} 人 / "
          f"キャリア {len(data['careers'])} 種 / 確率表 {len(data['tables'])} 行 / 統計 {len(data['stats'])} 件 / 絵 {len(data['assets'])} 枚 → {dst}")


if __name__ == '__main__':
    main()
